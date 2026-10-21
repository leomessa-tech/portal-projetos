import io
from datetime import datetime
from typing import List, Dict, Any, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.pdfgen import canvas

import logic
import i18n

REPORT_I18N = {
    "pt_BR": {
        "sheet_title": "Relatório de Projetos",
        "emitted_at": "Emitido em",
        "filters_applied": "Filtros aplicados",
        "total_listed": "Total de projetos listados",
        "col_id": "ID",
        "col_name": "Nome do Projeto",
        "col_prio": "Prioridade",
        "col_status_prazo": "Status (Prazo)",
        "col_status_manual": "Status Operacional",
        "col_start": "Data Início",
        "col_alert": "Data Aviso",
        "col_end": "Data Término",
        "col_days": "Dias Restantes",
        "col_prog": "Progresso",
        "col_resp": "Responsável",
        "col_email": "E-mail",
        "col_done": "Concluído",
        "col_obs": "Observações",
        "col_situation": "Situação",
        "col_resp_email": "Responsável / E-mail",
        "days_suffix": "dias",
        "days_delay": "dias (atraso)",
        "yes": "Sim",
        "no": "Não",
        "kpi_total": "Total",
        "page_x_of_y": "Página {page} de {total}",
        "days_left": "{days} d rest.",
    },
    "en_US": {
        "sheet_title": "Project Report",
        "emitted_at": "Emitted at",
        "filters_applied": "Applied filters",
        "total_listed": "Total projects listed",
        "col_id": "ID",
        "col_name": "Project Name",
        "col_prio": "Priority",
        "col_status_prazo": "Timeline Status",
        "col_status_manual": "Operational Status",
        "col_start": "Start Date",
        "col_alert": "Alert Date",
        "col_end": "End Date",
        "col_days": "Remaining Days",
        "col_prog": "Progress",
        "col_resp": "Responsible",
        "col_email": "Email",
        "col_done": "Completed",
        "col_obs": "Observations",
        "col_situation": "Timeline",
        "col_resp_email": "Responsible / Email",
        "days_suffix": "days",
        "days_delay": "days (overdue)",
        "yes": "Yes",
        "no": "No",
        "kpi_total": "Total",
        "page_x_of_y": "Page {page} of {total}",
        "days_left": "{days} d left",
    },
    "es_ES": {
        "sheet_title": "Informe de Proyectos",
        "emitted_at": "Emitido el",
        "filters_applied": "Filtros aplicados",
        "total_listed": "Total de proyectos listados",
        "col_id": "ID",
        "col_name": "Nombre del Proyecto",
        "col_prio": "Prioridad",
        "col_status_prazo": "Estado Plazos",
        "col_status_manual": "Estado Operativo",
        "col_start": "Fecha Inicio",
        "col_alert": "Fecha Aviso",
        "col_end": "Fecha Término",
        "col_days": "Días Restantes",
        "col_prog": "Progreso",
        "col_resp": "Responsable",
        "col_email": "Correo",
        "col_done": "Concluido",
        "col_obs": "Observaciones",
        "col_situation": "Situación",
        "col_resp_email": "Responsable / Correo",
        "days_suffix": "días",
        "days_delay": "días (retraso)",
        "yes": "Sí",
        "no": "No",
        "kpi_total": "Total",
        "page_x_of_y": "Página {page} de {total}",
        "days_left": "{days} d rest.",
    },
    "fr_FR": {
        "sheet_title": "Rapport de Projets",
        "emitted_at": "Émis le",
        "filters_applied": "Filtres appliqués",
        "total_listed": "Total des projets listés",
        "col_id": "ID",
        "col_name": "Nom du Projet",
        "col_prio": "Priorité",
        "col_status_prazo": "Statut Échéance",
        "col_status_manual": "Statut Opérationnel",
        "col_start": "Date Début",
        "col_alert": "Date Alerte",
        "col_end": "Date Fin",
        "col_days": "Jours Restants",
        "col_prog": "Progression",
        "col_resp": "Responsable",
        "col_email": "E-mail",
        "col_done": "Terminé",
        "col_obs": "Observations",
        "col_situation": "Situation",
        "col_resp_email": "Responsable / E-mail",
        "days_suffix": "jours",
        "days_delay": "jours (retard)",
        "yes": "Oui",
        "no": "Non",
        "kpi_total": "Total",
        "page_x_of_y": "Page {page} sur {total}",
        "days_left": "{days} j rest.",
    }
}

def _rt(key: str, lang: str = "pt_BR") -> str:
    dic = REPORT_I18N.get(lang, REPORT_I18N["pt_BR"])
    return dic.get(key, REPORT_I18N["pt_BR"].get(key, key))

def format_data_br(d_str: str) -> str:
    try:
        return logic.parse_data(d_str).strftime("%d/%m/%Y")
    except Exception:
        return str(d_str)

def traduzir_status_prazo(status_br: str, lang: str = "pt_BR") -> str:
    mapa = {
        "No Prazo": i18n.t("status_prazo_no_prazo", lang),
        "Chegando ao Fim": i18n.t("status_prazo_chegando", lang),
        "Atrasado": i18n.t("status_prazo_atrasado", lang),
        "Concluído": i18n.t("status_prazo_concluido", lang),
    }
    return mapa.get(status_br, status_br)

def traduzir_prioridade(prio_br: str, lang: str = "pt_BR") -> str:
    mapa = {
        "Crítica": i18n.t("prio_critica", lang),
        "Alta": i18n.t("prio_alta", lang),
        "Média": i18n.t("prio_media", lang),
        "Baixa": i18n.t("prio_baixa", lang),
    }
    return mapa.get(prio_br, prio_br)


# ==============================================================================
# GERADOR DE RELATÓRIO EXCEL (.XLSX)
# ==============================================================================

def gerar_relatorio_excel(
    projetos: List[Dict[str, Any]],
    lang: str = "pt_BR",
    filtros: Optional[Dict[str, str]] = None
) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = _rt("sheet_title", lang)
    ws.views.sheetView[0].showGridLines = True

    # Cores corporativas elegantes
    fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_title = Font(name="Segoe UI", size=16, bold=True, color="1E3A8A")
    font_sub = Font(name="Segoe UI", size=10, italic=True, color="475569")
    font_meta = Font(name="Segoe UI", size=10, color="334155")
    font_bold = Font(name="Segoe UI", size=10, bold=True, color="1E293B")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    # 1. Título do Relatório
    ws.merge_cells("A1:N1")
    ws["A1"] = i18n.t("portal_title", lang).replace("📋 ", "")
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 32

    # 2. Informações de Emissão e Filtros
    agora_fmt = datetime.now().strftime("%d/%m/%Y %H:%M")
    ws.merge_cells("A2:N2")
    filtros_txt = ", ".join([f"{k}: {v}" for k, v in (filtros or {}).items() if v and v not in ["Todos", "Todas", "All", "Tous"]])
    sub_txt = f"{_rt('emitted_at', lang)}: {agora_fmt}"
    if filtros_txt:
        sub_txt += f" | {_rt('filters_applied', lang)}: [{filtros_txt}]"
    sub_txt += f" | {_rt('total_listed', lang)}: {len(projetos)}"
    ws["A2"] = sub_txt
    ws["A2"].font = font_sub
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 20

    # Linha em branco para espaçamento visual
    ws.row_dimensions[3].height = 10

    # 3. Cabeçalho da Tabela
    colunas = [
        (_rt("col_id", lang), 8),
        (_rt("col_name", lang), 32),
        (_rt("col_prio", lang), 14),
        (_rt("col_status_prazo", lang), 18),
        (_rt("col_status_manual", lang), 20),
        (_rt("col_start", lang), 14),
        (_rt("col_alert", lang), 14),
        (_rt("col_end", lang), 14),
        (_rt("col_days", lang), 18),
        (_rt("col_prog", lang), 12),
        (_rt("col_resp", lang), 22),
        (_rt("col_email", lang), 26),
        (_rt("col_done", lang), 12),
        (_rt("col_obs", lang), 35)
    ]

    row_num = 4
    for col_idx, (col_nome, col_largura) in enumerate(colunas, 1):
        cell = ws.cell(row=row_num, column=col_idx, value=col_nome)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[row_num].height = 26

    # 4. Paleta de Cores de Status e Prioridade
    status_fills = {
        "Atrasado": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        "Chegando ao Fim": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
        "No Prazo": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
        "Concluído": PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    }
    status_fonts = {
        "Atrasado": Font(name="Segoe UI", size=10, bold=True, color="991B1B"),
        "Chegando ao Fim": Font(name="Segoe UI", size=10, bold=True, color="92400E"),
        "No Prazo": Font(name="Segoe UI", size=10, bold=True, color="166534"),
        "Concluído": Font(name="Segoe UI", size=10, bold=True, color="1E40AF")
    }

    prio_fills = {
        "Crítica": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        "Alta": PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid"),
        "Média": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
        "Baixa": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    }

    # 5. Preenchimento de Linhas de Projetos
    for p in projetos:
        row_num += 1
        ws.row_dimensions[row_num].height = 20
        
        status_orig = p.get("status", "No Prazo")
        prio_orig = p.get("prioridade", "Média")
        
        status_trad = traduzir_status_prazo(status_orig, lang)
        prio_trad = traduzir_prioridade(prio_orig, lang)
        
        dias_val = p.get("dias_restantes_termino", 0)
        dias_sufixo = _rt("days_suffix", lang)
        dias_atraso_sufixo = _rt("days_delay", lang)
        dias_str = f"{dias_val} {dias_sufixo}" if dias_val >= 0 else f"{abs(dias_val)} {dias_atraso_sufixo}"
        
        concluido_str = _rt("yes", lang) if p.get("concluido") else _rt("no", lang)
        
        row_data = [
            p.get("id"),
            p.get("nome"),
            prio_trad,
            status_trad,
            p.get("status_manual", "Em Execução"),
            format_data_br(p.get("data_inicio", "")),
            format_data_br(p.get("data_aviso", "")),
            format_data_br(p.get("data_termino", "")),
            dias_str,
            f"{p.get('progresso', 0)}%",
            p.get("responsavel"),
            p.get("email_responsavel"),
            concluido_str,
            p.get("observacoes", "")
        ]

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.font = font_meta
            cell.border = thin_border
            
            # Alinhamentos
            if col_idx in [1, 6, 7, 8, 9, 10, 13]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Estilo Status Prazo (coluna 4)
            if col_idx == 4 and status_orig in status_fills:
                cell.fill = status_fills[status_orig]
                cell.font = status_fonts[status_orig]

            # Estilo Prioridade (coluna 3)
            if col_idx == 3 and prio_orig in prio_fills:
                cell.fill = prio_fills[prio_orig]
                cell.font = font_bold

    # Ajuste das larguras de colunas
    for col_idx, (col_nome, largura_padrao) in enumerate(colunas, 1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = largura_padrao

    # Salvar em buffer de memória
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# ==============================================================================
# GERADOR DE RELATÓRIO PDF (.PDF)
# ==============================================================================

def make_numbered_canvas(lang: str = "pt_BR"):
    """Fábrica de Canvas numerado com internacionalização."""
    class LocalizedNumberedCanvas(canvas.Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._saved_page_states = []

        def showPage(self):
            self._saved_page_states.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            num_pages = len(self._saved_page_states)
            for state in self._saved_page_states:
                self.__dict__.update(state)
                self.draw_footer(num_pages)
                super().showPage()
            super().save()

        def draw_footer(self, page_count):
            self.saveState()
            self.setFont("Helvetica", 8)
            self.setFillColor(colors.HexColor("#64748B"))
            self.setStrokeColor(colors.HexColor("#CBD5E1"))
            self.setLineWidth(0.5)
            self.line(30, 28, 812, 28) # A4 Paisagem: largura total ~842 pt
            
            portal_label = i18n.t("page_title", lang)
            data_geracao = datetime.now().strftime("%d/%m/%Y %H:%M")
            self.drawString(30, 16, f"{portal_label} — {data_geracao}")
            
            pag_str = _rt("page_x_of_y", lang).format(page=self._pageNumber, total=page_count)
            self.drawRightString(812, 16, pag_str)
            self.restoreState()

    return LocalizedNumberedCanvas


def gerar_relatorio_pdf(
    projetos: List[Dict[str, Any]],
    lang: str = "pt_BR",
    filtros: Optional[Dict[str, str]] = None
) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=30,
        rightMargin=30,
        topMargin=28,
        bottomMargin=42
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1E3A8A'),
        spaceAfter=3
    )
    
    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#475569'),
        spaceAfter=10
    )

    th_style = ParagraphStyle(
        'TH',
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=1
    )

    cell_style = ParagraphStyle(
        'Cell',
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#1E293B')
    )

    cell_style_center = ParagraphStyle(
        'CellCenter',
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        alignment=1,
        textColor=colors.HexColor('#1E293B')
    )

    story = []

    # 1. Título do Documento
    titulo_texto = i18n.t("portal_title", lang).replace("📋 ", "")
    story.append(Paragraph(f"<b>{titulo_texto}</b>", title_style))

    # 2. Metadados de Emissão e Filtros
    agora_str = datetime.now().strftime("%d/%m/%Y %H:%M")
    filtros_str = ", ".join([f"{k}: {v}" for k, v in (filtros or {}).items() if v and v not in ["Todos", "Todas", "All", "Tous"]])
    sub_info = f"<b>{_rt('emitted_at', lang)}:</b> {agora_str}"
    if filtros_str:
        sub_info += f" &nbsp;|&nbsp; <b>{_rt('filters_applied', lang)}:</b> {filtros_str}"
    sub_info += f" &nbsp;|&nbsp; <b>{_rt('total_listed', lang)}:</b> {len(projetos)}"
    story.append(Paragraph(sub_info, meta_style))

    # 3. Resumo de Métricas (Banner de KPIs)
    total_p = len(projetos)
    n_prazo = sum(1 for p in projetos if p.get("status") == "No Prazo")
    n_aviso = sum(1 for p in projetos if p.get("status") == "Chegando ao Fim")
    n_atraso = sum(1 for p in projetos if p.get("status") == "Atrasado")
    n_conc = sum(1 for p in projetos if p.get("status") == "Concluído")

    kpi_data = [[
        Paragraph(f"<b>{_rt('kpi_total', lang)}:</b> {total_p}", cell_style_center),
        Paragraph(f"<b>{i18n.t('kpi_no_prazo', lang)}:</b> {n_prazo}", cell_style_center),
        Paragraph(f"<b>{i18n.t('kpi_chegando_fim', lang)}:</b> {n_aviso}", cell_style_center),
        Paragraph(f"<b>{i18n.t('kpi_atrasados', lang)}:</b> {n_atraso}", cell_style_center),
        Paragraph(f"<b>{i18n.t('kpi_concluidos', lang)}:</b> {n_conc}", cell_style_center),
    ]]
    kpi_table = Table(kpi_data, colWidths=[156, 156, 156, 156, 156])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 10))

    # 4. Tabela de Projetos
    # Largura útil da folha A4 em modo paisagem: ~782 pt
    col_widths = [160, 65, 80, 85, 60, 60, 60, 65, 127]
    
    headers = [
        Paragraph(f"<b>{_rt('col_name', lang)}</b>", th_style),
        Paragraph(f"<b>{_rt('col_prio', lang)}</b>", th_style),
        Paragraph(f"<b>{_rt('col_status_prazo', lang)}</b>", th_style),
        Paragraph(f"<b>{_rt('col_status_manual', lang)}</b>", th_style),
        Paragraph(f"<b>{_rt('col_start', lang)}</b>", th_style),
        Paragraph(f"<b>{_rt('col_alert', lang)}</b>", th_style),
        Paragraph(f"<b>{_rt('col_end', lang)}</b>", th_style),
        Paragraph(f"<b>{_rt('col_situation', lang)}</b>", th_style),
        Paragraph(f"<b>{_rt('col_resp_email', lang)}</b>", th_style)
    ]
    
    table_data = [headers]

    for p in projetos:
        status_orig = p.get("status", "No Prazo")
        dias_val = p.get("dias_restantes_termino", 0)
        prio_orig = p.get("prioridade", "Média")
        
        status_trad = traduzir_status_prazo(status_orig, lang)
        prio_trad = traduzir_prioridade(prio_orig, lang)
        
        if status_orig == "Concluído":
            sit_txt = f"<font color='#1E40AF'><b>{status_trad}</b></font>"
        elif status_orig == "Atrasado":
            sit_txt = f"<font color='#991B1B'><b>-{abs(dias_val)} d</b></font>"
        elif status_orig == "Chegando ao Fim":
            txt_d = _rt("days_left", lang).format(days=dias_val)
            sit_txt = f"<font color='#92400E'><b>{txt_d}</b></font>"
        else:
            txt_d = _rt("days_left", lang).format(days=dias_val)
            sit_txt = f"<font color='#166534'><b>{txt_d}</b></font>"

        row = [
            Paragraph(f"<b>{p.get('nome')}</b>", cell_style),
            Paragraph(prio_trad, cell_style_center),
            Paragraph(status_trad, cell_style_center),
            Paragraph(p.get("status_manual", "Em Execução"), cell_style_center),
            Paragraph(format_data_br(p.get("data_inicio", "")), cell_style_center),
            Paragraph(format_data_br(p.get("data_aviso", "")), cell_style_center),
            Paragraph(format_data_br(p.get("data_termino", "")), cell_style_center),
            Paragraph(sit_txt, cell_style_center),
            Paragraph(f"{p.get('responsavel')}<br/><font color='#64748B'>{p.get('email_responsavel')}</font>", cell_style)
        ]
        table_data.append(row)

    main_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    
    t_style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
    ]

    # Cores alternadas para as linhas de dados
    for r in range(1, len(table_data)):
        if r % 2 == 0:
            t_style.append(('BACKGROUND', (0, r), (-1, r), colors.HexColor('#F8FAFC')))
        else:
            t_style.append(('BACKGROUND', (0, r), (-1, r), colors.white))

    main_table.setStyle(TableStyle(t_style))
    story.append(main_table)

    # Compilação do PDF
    canvas_cls = make_numbered_canvas(lang)
    doc.build(story, canvasmaker=canvas_cls)
    buffer.seek(0)
    return buffer.getvalue()


# ==============================================================================
# GERADOR DE PROPOSTA COMERCIAL & TÉCNICA EM PDF
# ==============================================================================

def gerar_proposta_pdf(prop: Dict[str, Any]) -> bytes:
    from reportlab.lib.pagesizes import A4
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    
    style_normal = ParagraphStyle('PropNormal', parent=styles['Normal'], fontName='Helvetica', fontSize=10, leading=14, textColor=colors.HexColor('#1E293B'))
    style_bold = ParagraphStyle('PropBold', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10, leading=14, textColor=colors.HexColor('#0F172A'))
    style_title = ParagraphStyle('PropTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=12, leading=16, textColor=colors.HexColor('#1E3A8A'))
    style_th = ParagraphStyle('PropTH', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10, leading=13, textColor=colors.white)
    style_small = ParagraphStyle('PropSmall', parent=styles['Normal'], fontName='Helvetica', fontSize=8.5, leading=12, textColor=colors.HexColor('#475569'))
    
    story = []

    ref = prop.get("referencia", "26-001")
    data_prop = prop.get("data_proposta", "") or logic.formatar_data_extenso_pt()
    hotel = prop.get("hotel_cliente", "Hotel Cliente")
    chamado = prop.get("chamado_servicenow", "N/A") or "N/A"
    contato_ac = prop.get("contato_ac", "Hotel") or "Hotel"
    descricao = prop.get("descricao_servico", "Instalação e Configuração de Interface")
    valor = logic.formatar_moeda_br(prop.get("valor_total", 0.0))
    condicoes = prop.get("condicoes_comerciais", logic.CONDICOES_COMERCIAIS_PADRAO) or logic.CONDICOES_COMERCIAIS_PADRAO
    
    razao = prop.get("razao_social", "") or "________________________________________________"
    cnpj = prop.get("cnpj", "") or "_______________________________________"
    nome_ass = prop.get("nome_assinatura", "") or "_______________________________________"
    cargo_ass = prop.get("cargo_assinatura", "") or "_______________________________________"
    email_cob = prop.get("email_cobranca", "") or "___________________________________________@accor.com"

    # 1. Tabela de Cabeçalho (Ref e Data)
    header_data = [
        [
            Paragraph(f"<b>Ref: {ref}</b>", style_title),
            Paragraph(f"<b>{data_prop}</b>", ParagraphStyle('PropDate', parent=style_normal, alignment=2, fontName='Helvetica-Bold', textColor=colors.HexColor('#1E3A8A')))
        ]
    ]
    header_table = Table(header_data, colWidths=[260, 263])
    header_table.setStyle(TableStyle([
        ('BOX', (0,0), (-1,-1), 1.5, colors.HexColor('#1E3A8A')),
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
        ('TOPPADDING', (0,0), (-1,-1), 8),
        ('BOTTOMPADDING', (0,0), (-1,-1), 8),
        ('LEFTPADDING', (0,0), (-1,-1), 10),
        ('RIGHTPADDING', (0,0), (-1,-1), 10),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 15))

    # 2. Informações do Hotel e Chamado
    story.append(Paragraph(f"<b>{hotel}</b>", ParagraphStyle('HotelH', parent=style_normal, fontSize=13, fontName='Helvetica-Bold', textColor=colors.HexColor('#0F172A'), spaceAfter=4)))
    story.append(Paragraph(f"<b>Chamado Service Now:</b> {chamado}", style_normal))
    story.append(Paragraph(f"<b>Contato:</b> {contato_ac}", style_normal))
    story.append(Spacer(1, 14))

    # 3. Introdução
    story.append(Paragraph("Conforme a vossa solicitação, apresentamos nossa proposta comercial e técnica para os seguintes itens:", style_normal))
    story.append(Spacer(1, 10))

    # 4. Tabela de Itens e Preços
    items_data = [
        [Paragraph("<b>Descrição do Serviço</b>", style_th), Paragraph("<b>Total</b>", ParagraphStyle('PropTHRight', parent=style_th, alignment=2))],
        [Paragraph(descricao, style_normal), Paragraph(f"<b>{valor}</b>", ParagraphStyle('PropValRight', parent=style_normal, alignment=2, fontName='Helvetica-Bold', textColor=colors.HexColor('#1E3A8A')))]
    ]
    items_table = Table(items_data, colWidths=[383, 140])
    items_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E3A8A')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0,0), (-1,-1), 7),
        ('BOTTOMPADDING', (0,0), (-1,-1), 7),
        ('LEFTPADDING', (0,0), (-1,-1), 8),
        ('RIGHTPADDING', (0,0), (-1,-1), 8),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(items_table)
    story.append(Spacer(1, 16))

    # 5. Condições Comerciais
    story.append(Paragraph("<b>Condições comerciais:</b>", style_bold))
    story.append(Spacer(1, 6))
    for linha in condicoes.splitlines():
        if linha.strip():
            story.append(Paragraph(f"• {linha.strip()}", style_small))
            story.append(Spacer(1, 2))
    story.append(Spacer(1, 14))

    # 6. Quadro "De acordo"
    story.append(Paragraph("<b>De acordo:</b>", style_bold))
    story.append(Spacer(1, 6))
    
    de_acordo_data = [
        [Paragraph(f"<b>Razão Social:</b> {razao}", style_small), Paragraph(f"<b>CNPJ:</b> {cnpj}", style_small)],
        [Paragraph(f"<b>NOME:</b> {nome_ass}", style_small), Paragraph(f"<b>CARGO:</b> {cargo_ass}", style_small)],
        [Paragraph(f"<b>Email de cobrança:</b> {email_cob}", style_small), Paragraph("", style_small)],
        [Paragraph("<b>Assinatura:</b> ____________________________________________", style_small), Paragraph("<b>Data Assinatura:</b> _______/________/________", style_small)]
    ]
    de_acordo_table = Table(de_acordo_data, colWidths=[280, 243])
    de_acordo_table.setStyle(TableStyle([
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#94A3B8')),
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('LEFTPADDING', (0,0), (-1,-1), 8),
        ('RIGHTPADDING', (0,0), (-1,-1), 8),
        ('SPAN', (0, 2), (1, 2)), # Email ocupa as duas colunas
    ]))
    story.append(de_acordo_table)

    # Compilação
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()
